"""Exportação da lista de prioridade de agenda (IIP) em Excel formatado."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .avaliador_afinidade import carregar_agenda
from .composicao import DATA_DIR
from .interesse import FAIXAS, PESOS_COMPONENTES, AnalisadorInteresse
from .painel import tabela, todos
from .schemas import Composicao

AZUL, AZUL_CLARO = "1B2A4A", "EAF1F8"
COR_FAIXA = {"A": "27AE60", "B": "7DBE5A", "C": "F1C40F", "D": "D5D8DC"}


def _cabecalho(ws, linha: int) -> None:
    for cel in ws[linha]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=AZUL)
        cel.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)


def exportar_prioridade(comp: Composicao, top: int | None = None, caminho: Path | None = None) -> Path:
    df = tabela(comp)
    df = df[df["iip"] > 0].sort_values(["iip", "n_proposicoes", "votos"], ascending=False)
    if top:
        df = df.head(top)
    df = df.reset_index(drop=True)
    df.insert(0, "Posição", range(1, len(df) + 1))

    principal = df[[
        "Posição", "prioridade", "iip", "casa", "uf", "nome", "partido", "trajetoria",
        "iip_autoria", "iip_relatoria", "iip_comissoes", "iip_frentes",
        "n_proposicoes", "n_relatorias", "orgaos_tech", "frentes_tech", "votos", "nome_completo",
    ]].rename(columns={
        "prioridade": "Prioridade", "iip": "IIP (0-100)", "casa": "Casa (2027)", "uf": "UF",
        "nome": "Parlamentar", "partido": "Partido", "trajetoria": "Trajetória",
        "iip_autoria": "Autoria", "iip_relatoria": "Relatoria (Senado)", "iip_comissoes": "Comissões",
        "iip_frentes": "Frentes", "n_proposicoes": "Nº proposições", "n_relatorias": "Nº relatorias",
        "orgaos_tech": "Comissões/órgãos temáticos", "frentes_tech": "Frentes parlamentares",
        "votos": "Votos 2026", "nome_completo": "Nome completo",
    })
    principal["Parlamentar"] = principal["Parlamentar"].str.title()
    principal["Nome completo"] = principal["Nome completo"].fillna("").str.title()

    # Detalhe de evidências, uma linha por sinal
    por_chave = {p.chave: p for p in todos(comp)}
    nomes_temas = {t["id"]: t["nome"] for t in carregar_agenda()["temas"]}
    evid = []
    for _, r in df.iterrows():
        i = AnalisadorInteresse.carregar(por_chave[r["chave"]])
        if not i:
            continue
        base = {"Posição": r["Posição"], "Parlamentar": r["nome"].title(), "Partido/UF": f"{r['partido']}/{r['uf']}"}
        for tipo, itens in [("Autoria", i.proposicoes), ("Relatoria", i.relatorias),
                            ("Comissão/órgão", i.orgaos), ("Frente parlamentar", i.frentes)]:
            for it in itens:
                temas = ""
                if "[" in it:
                    ids = it.split("[", 1)[1].split("]", 1)[0].split(", ")
                    temas = ", ".join(nomes_temas.get(t, t) for t in ids)
                evid.append({**base, "Tipo de sinal": tipo, "Temas": temas, "Detalhe": it})

    pesos_txt = ", ".join(f"{k} {v:.0%}" for k, v in PESOS_COMPONENTES.items())
    faixas_txt = "; ".join(f"{r} (IIP ≥ {c})" for c, r in FAIXAS) + "; D — Interesse baixo (IIP > 0)"
    metodo = pd.DataFrame({"Índice de Interesse na Pauta ABES (IIP)": [
        "Universo: membros do Congresso 2027-2031 que exercem mandato na legislatura atual — deputados reeleitos, "
        "senadores reeleitos ou com mandato até 2031, e quem troca de casa (Câmara ↔ Senado). Novatos ficam de fora "
        "por não terem histórico mensurável.",
        "Período: 01/02/2023 até hoje. Fontes: APIs de Dados Abertos da Câmara e do Senado.",
        "Autoria: PL/PLP/PEC cujas ementas tratam dos temas da agenda (palavras-chave em eleicoes/agenda_abes.json), "
        "ponderados pelo peso de cada tema; satura em 3 proposições por tema. No Senado, coautoria vale metade.",
        "Relatoria: matérias da agenda relatadas no Senado (satura em 4). A Câmara não publica relatorias por deputado; "
        "para deputados o componente é desconsiderado e os demais pesos são renormalizados.",
        "Comissões: participação em CCTI/CCOM (Câmara), CCT/CCDD (Senado) e em comissões especiais, grupos de trabalho "
        "ou subcomissões sobre temas da agenda. Presidente 100%, vice/relator/coordenador 80%, titular 60%, "
        "suplente 30%, +15% por órgão adicional.",
        "Frentes: frentes parlamentares ligadas à pauta digital (satura em 3).",
        f"Pesos: {pesos_txt}. Faixas: {faixas_txt}.",
        "LIMITE: o IIP mede INTERESSE/ENGAJAMENTO na pauta, não ALINHAMENTO às posições da ABES. "
        "Um parlamentar muito engajado pode defender posições contrárias — a qualificação de posição vem da análise com IA.",
        f"Composição coletada em {comp.data_coleta} (TSE — resultado oficial; Senado — API de Dados Abertos).",
    ]})

    caminho = caminho or DATA_DIR / (f"prioridade_agenda_ABES_top{top}.xlsx" if top else "prioridade_agenda_ABES.xlsx")
    fino = Side(style="thin", color="D5D8DC")
    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        principal.to_excel(w, sheet_name="Prioridade de agenda", index=False, startrow=2)
        pd.DataFrame(evid).to_excel(w, sheet_name="Evidências", index=False)
        metodo.to_excel(w, sheet_name="Metodologia", index=False)

        ws = w.sheets["Prioridade de agenda"]
        ws["A1"] = "Prioridade de agenda ABES — Congresso 2027-2031 (parlamentares com histórico de interesse na pauta)"
        ws["A1"].font = Font(bold=True, size=14, color=AZUL)
        ws["A2"] = "IIP mede interesse/engajamento na pauta, não alinhamento — ver aba Metodologia."
        ws["A2"].font = Font(italic=True, color="C0392B")
        _cabecalho(ws, 3)
        larguras = {"Posição": 8, "Prioridade": 20, "IIP (0-100)": 9, "Casa (2027)": 9, "UF": 5, "Parlamentar": 26,
                    "Partido": 13, "Trajetória": 22, "Nº proposições": 11, "Nº relatorias": 10,
                    "Comissões/órgãos temáticos": 50, "Frentes parlamentares": 50, "Votos 2026": 11,
                    "Nome completo": 34}
        pct = {"Autoria", "Relatoria (Senado)", "Comissões", "Frentes"}
        for j, col in enumerate(principal.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = larguras.get(col, 10)
            for i in range(4, 4 + len(principal)):
                cel = ws.cell(row=i, column=j)
                cel.border = Border(bottom=fino)
                cel.alignment = Alignment(vertical="top", wrap_text=col in ("Comissões/órgãos temáticos",
                                                                             "Frentes parlamentares"))
                if col in pct:
                    cel.number_format = "0%"
                if col == "Votos 2026":
                    cel.number_format = "#,##0"
                if col == "Prioridade" and cel.value:
                    cel.fill = PatternFill("solid", fgColor=COR_FAIXA.get(str(cel.value)[0], "FFFFFF"))
                    cel.font = Font(bold=True)
                elif i % 2 == 0:
                    cel.fill = PatternFill("solid", fgColor=AZUL_CLARO)
        ws.row_dimensions[3].height = 45
        ws.auto_filter.ref = f"A3:{get_column_letter(len(principal.columns))}{3 + len(principal)}"

        we = w.sheets["Evidências"]
        _cabecalho(we, 1)
        for j, wd in enumerate([8, 26, 16, 18, 40, 120], 1):
            we.column_dimensions[get_column_letter(j)].width = wd
        we.auto_filter.ref = we.dimensions
        wm = w.sheets["Metodologia"]
        _cabecalho(wm, 1)
        wm.column_dimensions["A"].width = 160
        for row in wm.iter_rows(min_row=2):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")
    return caminho

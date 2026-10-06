"""Relatório ABES — Prioridade de agenda no Congresso 2027-2031 (Índice de Interesse na Pauta)."""

from __future__ import annotations

import tempfile
from collections import Counter
from datetime import date
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from eleicoes import composicao as comp_mod
from eleicoes.avaliador_afinidade import carregar_agenda
from eleicoes.composicao import DATA_DIR
from eleicoes.interesse import FAIXAS, PESOS_COMPONENTES, AnalisadorInteresse
from eleicoes.painel import tabela, todos
from relatorio.gerador import (
    ASSETS, BRANCO, CINZA_BG, CINZA_LINHA, CINZA_RODAPE, DOURADO, PRETO,
    _build_header, _cell_text, _heading1, _heading2, _page_setup, _para,
    _set_cell_bg, _set_font,
)

COR_FAIXA = {"A": "2E7D32", "B": "7DBE5A", "C": "D9A61A", "D": "BDBDBD"}
MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
         "agosto", "setembro", "outubro", "novembro", "dezembro"]


def _data_extenso(d: date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def _rodape(doc) -> None:
    footer = doc.sections[0].footer
    for p in footer.paragraphs:
        p.clear()
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "6"), ("w:color", CINZA_LINHA)):
        top.set(qn(k), v)
    pBdr.append(top)
    pPr.append(pBdr)
    _set_font(p.add_run("Avenida Ibirapuera, 2.907 – 8º andar, cj. 811 – CEP 04029200 – São Paulo – SP   |  "
                        "tel: (11) 5094-3100"), 7.5, color=CINZA_RODAPE)
    ps = footer.add_paragraph()
    ps.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    social = ASSETS / "abes-footer-social.png"
    if social.exists():
        ps.add_run().add_picture(str(social), width=Cm(5.5))
    pid = footer.add_paragraph()
    pid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(pid.add_run(f"ABES · Diretoria de RIG · Prioridade de agenda – Congresso 2027-2031 · "
                          f"{date.today().strftime('%d/%m/%Y')}"), 7.5, color=CINZA_RODAPE)


_ORDEM_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
              "rPr", "sectPr", "pPrChange"]
_DEPOIS_DE_PBDR = {"shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
                   "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
                   "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
                   "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
                   "rPr", "sectPr", "pPrChange"}


def _ordenar_xml(doc) -> None:
    """Ajusta a ordem exigida pelo schema OOXML (pBdr dentro de pPr; zoom com percent)."""
    partes = [doc.element.body] + [s.footer._element for s in doc.sections] + [s.header._element for s in doc.sections]
    for raiz in partes:
        for pPr in raiz.iter(qn("w:pPr")):
            filhos = list(pPr)
            pos = {n: i for i, n in enumerate(_ORDEM_PPR)}
            filhos.sort(key=lambda c: pos.get(c.tag.split("}")[1], len(_ORDEM_PPR)))
            for c in filhos:
                pPr.remove(c)
                pPr.append(c)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")


def _linha_nao_quebra(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def _repete_cabecalho(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trPr.append(el)


def _tabela(doc, cabecalho: list[str], linhas: list[list[str]], larguras_cm: list[float],
            tamanho=8.5, cores_col0: list[str] | None = None, centro: set[int] | None = None):
    centro = centro or set()
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0]
    _linha_nao_quebra(hdr)
    _repete_cabecalho(hdr)
    for j, h in enumerate(cabecalho):
        _set_cell_bg(hdr.cells[j], PRETO)
        _cell_text(hdr.cells[j], h, size=tamanho, bold=True, color=BRANCO,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, lin in enumerate(linhas):
        row = t.add_row()
        _linha_nao_quebra(row)
        for j, v in enumerate(lin):
            cel = row.cells[j]
            if j == 0 and cores_col0:
                _set_cell_bg(cel, cores_col0[i])
                _cell_text(cel, v, size=tamanho, bold=True, color=BRANCO, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                if i % 2:
                    _set_cell_bg(cel, CINZA_BG)
                _cell_text(cel, v, size=tamanho,
                           align=WD_ALIGN_PARAGRAPH.CENTER if j in centro else WD_ALIGN_PARAGRAPH.LEFT)
    for row in t.rows:
        for j, w in enumerate(larguras_cm):
            row.cells[j].width = Cm(w)
    return t


def _bullet(doc, texto: str, negrito: str = "") -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if negrito:
        _set_font(p.add_run(negrito), 10.5, bold=True)
    _set_font(p.add_run(texto), 10.5)


def _grafico_faixas(h: pd.DataFrame) -> Path:
    ordem = ["A — Prioridade alta", "B — Prioridade média", "C — Interesse pontual", "D — Interesse baixo"]
    casas = ["Câmara", "Senado"]
    fig, ax = plt.subplots(figsize=(8.2, 2.6), dpi=200)
    esquerda = [0, 0]
    for faixa in ordem:
        vals = [int(((h["casa"] == c) & (h["prioridade"] == faixa)).sum()) for c in casas]
        barras = ax.barh(casas, vals, left=esquerda, color="#" + COR_FAIXA[faixa[0]], label=faixa,
                         edgecolor="white", height=0.55)
        for b, v in zip(barras, vals):
            if v == 0:
                continue
            x = b.get_x() + b.get_width() / 2
            if v >= 6:  # cabe dentro da barra
                ax.text(x, b.get_y() + b.get_height() / 2, str(v),
                        ha="center", va="center", fontsize=8, color="white", fontweight="bold")
            else:       # segmento estreito: rótulo acima da barra, para nenhum valor ficar oculto
                ax.annotate(str(v), (x, b.get_y()), xytext=(0, 3), textcoords="offset points",
                            ha="center", va="bottom", fontsize=8.5, fontweight="bold",
                            color="#" + COR_FAIXA[faixa[0]])
        esquerda = [e + v for e, v in zip(esquerda, vals)]
    ax.invert_yaxis()
    ax.set_xlabel("Parlamentares com histórico de mandato", fontsize=8, color="#404040")
    ax.tick_params(labelsize=9)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(ncol=4, fontsize=7.5, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.22))
    fig.tight_layout()
    caminho = Path(tempfile.mktemp(suffix=".png"))
    fig.savefig(caminho, bbox_inches="tight")
    plt.close(fig)
    return caminho


def _sinais_resumo(i, nomes_temas: dict[str, str], max_temas=3) -> str:
    temas = sorted(i.autoria_por_tema.items(), key=lambda kv: -kv[1])[:max_temas]
    partes = []
    if temas:
        partes.append("Autoria: " + ", ".join(f"{nomes_temas[t]} ({n:g})" for t, n in temas))
    if i.relatorias:
        partes.append(f"Relatorias: {len(i.relatorias)}")
    orgaos = sorted({o.split(" — ")[0] for o in i.orgaos})
    if orgaos:
        partes.append("Órgãos: " + ", ".join(orgaos[:5]) + ("…" if len(orgaos) > 5 else ""))
    return " · ".join(partes) or "—"


def gerar(caminho: Path | None = None) -> Path:
    comp = comp_mod.carregar()
    if comp is None:
        raise RuntimeError("Composição não coletada — rode: python eleicoes_main.py composicao")
    agenda = carregar_agenda()
    nomes_temas = {t["id"]: t["nome"] for t in agenda["temas"]}
    df = tabela(comp)
    por_chave = {p.chave: p for p in todos(comp)}
    hist = df[df["prioridade"] != "Sem histórico (novo no Congresso)"]
    h = hist[hist["iip"] > 0].sort_values(["iip", "n_proposicoes"], ascending=False)
    interesses = {k: AnalisadorInteresse.carregar(por_chave[k]) for k in h["chave"]}
    fa = h[h["prioridade"].str.startswith("A")]
    fb = h[h["prioridade"].str.startswith("B")]
    novos = df[df["trajetoria"] == "Novo(a)"]
    hoje = date.today()

    doc = Document()
    _page_setup(doc)
    _build_header(doc)
    _rodape(doc)
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(10.5)

    # ── Capa ────────────────────────────────────────────────────────────
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    esq, dir_ = t.rows[0].cells
    esq.width, dir_.width = Cm(11.5), Cm(4.5)
    _set_cell_bg(esq, PRETO)
    _set_cell_bg(dir_, "2E7D32")
    p = esq.paragraphs[0]
    p.paragraph_format.space_before = Pt(10)
    _set_font(p.add_run("RELATÓRIO DE PROSPECÇÃO POLÍTICA"), 9, bold=True, color=DOURADO)
    p = esq.add_paragraph()
    _set_font(p.add_run("Prioridade de agenda ABES no Congresso Nacional – Legislatura 2027-2031"),
              13, bold=True, color=BRANCO)
    p = esq.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    _set_font(p.add_run(f"Parlamentares eleitos em 04/10/2026 e senadores com mandato até 2031, "
                        f"hierarquizados pelo interesse prévio na pauta de software e tecnologia · "
                        f"{_data_extenso(hoje)}"), 9, color="CCCCCC")
    for k, (txt, tam) in enumerate([(str(len(h)), 26), ("parlamentares com interesse prévio na pauta", 8),
                                    (f"{len(fa) + len(fb)}", 18), ("em prioridade alta ou média", 8)]):
        pp = dir_.paragraphs[0] if k == 0 else dir_.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if k == 0:
            pp.paragraph_format.space_before = Pt(8)
        _set_font(pp.add_run(txt), tam, bold=tam > 10, color=BRANCO)
    doc.add_paragraph()

    # ── 1. Síntese executiva ────────────────────────────────────────────
    _heading1(doc, "1. Síntese executiva")
    _para(doc, (
        f"A partir de 1º de fevereiro de 2027, {len(hist)} dos 594 parlamentares federais "
        f"({len(hist) / 594:.0%}) terão exercido mandato na legislatura atual: são deputados reeleitos, senadores "
        f"que permanecem ou foram reeleitos e parlamentares que trocam de casa. Para esse grupo há histórico "
        f"mensurável, e {len(h)} demonstraram algum interesse na pauta da ABES entre 2023 e 2026 — por autoria de "
        f"proposições, relatorias, participação em comissões temáticas ou em frentes parlamentares."
    ))
    _para(doc, (
        f"O Índice de Interesse na Pauta ABES (IIP) hierarquiza esse grupo em quatro faixas. {len(fa)} parlamentares "
        f"estão na faixa A (prioridade alta) e {len(fb)} na faixa B (prioridade média): são os interlocutores "
        f"naturais para a agenda institucional da ABES no início da nova legislatura. Os {len(novos)} parlamentares "
        f"sem mandato anterior não são avaliados por este índice e demandam mapeamento próprio, baseado em "
        f"propostas de campanha."
    ))
    _para(doc, (
        "Ressalva central: o IIP mede interesse e engajamento na pauta, não alinhamento às posições da ABES. "
        "Um parlamentar muito ativo em temas digitais pode defender posições contrárias às do setor — a "
        "qualificação de posição (nota de −2 a +2 por tema) é a etapa seguinte recomendada para as faixas A e B."
    ), bold=True, size=10)

    # ── 2. Nova composição ──────────────────────────────────────────────
    _heading1(doc, "2. Nova composição do Congresso")
    traj = Counter(df["trajetoria"])
    camara_novos = int(((df["casa"] == "Câmara") & (df["trajetoria"] == "Novo(a)")).sum())
    senado_novos = int(((df["casa"] == "Senado") & (df["trajetoria"] == "Novo(a)")).sum())
    _tabela(doc, ["Casa", "Total", "Com mandato na legislatura atual", "Sem mandato anterior"], [
        ["Câmara dos Deputados", "513",
         f"{traj['Reeleito(a) deputado(a)']} reeleitos + {traj['Senado → Câmara']} vindos do Senado",
         str(camara_novos)],
        ["Senado Federal", "81",
         f"{traj['Permanece (mandato até 2031)']} com mandato até 2031 + {traj['Reeleito(a) senador(a)']} reeleitos + "
         f"{traj['Câmara → Senado']} vindos da Câmara", str(senado_novos)],
    ], [4.2, 1.6, 7.4, 2.8], tamanho=9, centro={1, 3})
    _para(doc, "")
    _para(doc, (
        f"Fontes: resultado oficial do TSE para o 1º turno (100% das seções totalizadas em todas as UFs"
        f"{'; ver alertas abaixo' if comp.alertas else ''}; coleta em {comp.data_coleta}) e API de Dados Abertos do Senado. Cada UF elegeu dois senadores em 2026; "
        f"o terceiro senador de cada estado cumpre mandato iniciado em 2023."
    ), size=9, color=CINZA_RODAPE)
    if comp.alertas:
        _heading2(doc, "Alertas sobre os dados do TSE")
        for al in comp.alertas:
            _bullet(doc, f"{al}.")
    alertas = [p for p in comp.senadores if any("governo" in a for a in p.alertas)]
    if alertas:
        _heading2(doc, "Alertas sobre a composição do Senado")
        for p in alertas:
            al = next(a for a in p.alertas if "governo" in a)
            _bullet(doc, f" ({p.partido}/{p.uf}) — {al}.", negrito=p.nome_urna)

    # ── 3. Metodologia ──────────────────────────────────────────────────
    _heading1(doc, "3. Metodologia — Índice de Interesse na Pauta ABES (IIP)")
    _para(doc, (
        "O IIP (0 a 100) combina quatro sinais públicos, coletados nas APIs de Dados Abertos da Câmara e do "
        "Senado para o período de 01/02/2023 até a data deste relatório. Os temas e palavras-chave correspondem "
        "à agenda ABES (inteligência artificial, dados, tributação de software, P&D, cibersegurança, "
        "infraestrutura digital, governo digital, talentos, propriedade intelectual, plataformas e startups)."
    ))
    _tabela(doc, ["Sinal", "Peso", "O que mede"], [
        ["Autoria", f"{PESOS_COMPONENTES['autoria']:.0%}",
         "PL, PLP e PEC nos temas da agenda, ponderados pelo peso de cada tema (satura em 3 por tema). "
         "No Senado, coautoria vale metade."],
        ["Relatoria", f"{PESOS_COMPONENTES['relatoria']:.0%}",
         "Matérias da agenda relatadas no Senado. A Câmara não publica relatorias por deputado: para "
         "deputados o componente é desconsiderado e os demais pesos são renormalizados."],
        ["Comissões", f"{PESOS_COMPONENTES['comissoes']:.0%}",
         "CCTI e CCOM (Câmara), CCT e CCDD (Senado), comissões especiais e grupos de trabalho sobre temas da "
         "agenda. Presidente > vice/relator > titular > suplente."],
        ["Frentes", f"{PESOS_COMPONENTES['frentes']:.0%}",
         "Frentes parlamentares ligadas à pauta (Frente Digital, Desoneração da Folha, Propriedade Intelectual, "
         "Cibersegurança, entre outras). Sinal fraco: frentes mistas têm mais de 100 signatários."],
    ], [2.6, 1.4, 12.0], tamanho=8.5, centro={1})
    _para(doc, "")
    _heading2(doc, "Temas da agenda ABES")
    temas_ag = agenda["temas"]
    peso_total = sum(t.get("peso", 1) for t in temas_ag)
    por_peso = {w: [t["nome"].lower() for t in temas_ag if t.get("peso", 1) == w] for w in (3, 2, 1)}
    _para(doc, (
        f"A categorização parte de {len(temas_ag)} temas que traduzem a agenda de políticas públicas defendida "
        "pela ABES para o setor de software e serviços de tecnologia. Cada tema reúne uma descrição de escopo, "
        "a posição institucional da ABES e um conjunto de palavras-chave. Essas palavras-chave são procuradas nas "
        "ementas das proposições — sem considerar acentos e sempre como palavra inteira — para identificar a que "
        "tema cada projeto pertence; um mesmo projeto pode se enquadrar em mais de um tema. Os nomes de comissões "
        "e de frentes parlamentares passam pelo mesmo filtro, o que permite reconhecer órgãos dedicados a esses "
        "assuntos."
    ))
    _para(doc, (
        f"Cada tema recebe um peso de 1 a 3, que expressa sua prioridade relativa na agenda (soma total: "
        f"{peso_total}). Têm peso 3 {' e '.join(por_peso[3])}, que concentram as matérias de maior impacto para o "
        f"setor na próxima legislatura; têm peso 2 {', '.join(por_peso[2][:-1])} e {por_peso[2][-1]}; e têm peso "
        f"1 {', '.join(por_peso[1][:-1])} e {por_peso[1][-1]}. Na prática, atuar em temas de peso maior eleva mais "
        "o componente de autoria do IIP. Vale lembrar que o enquadramento num tema indica que o parlamentar atua "
        "no assunto, e não em que direção: um projeto restritivo e um projeto favorável à posição da ABES contam "
        "da mesma forma. Temas, pesos e palavras-chave são uma versão preliminar, sujeita à validação da "
        "Diretoria de RIG da ABES."
    ))
    linhas = [[t["nome"], str(t.get("peso", 1)), t.get("descricao", ""),
               ", ".join(t["palavras_chave"][:4])] for t in temas_ag]
    _tabela(doc, ["Tema", "Peso", "O que abrange", "Exemplos de palavras-chave (grafadas sem acento, como são comparadas)"], linhas,
            [3.4, 1.2, 6.8, 4.6], tamanho=8, centro={1})
    _para(doc, "")

    _heading2(doc, "Faixas de prioridade")
    _para(doc, "O IIP classifica os parlamentares em quatro faixas, que orientam a intensidade e o formato do "
               "relacionamento institucional da ABES ao longo da legislatura.")
    cortes = {r[0]: c for c, r in FAIXAS}
    descricoes = {
        "A": ("Prioridade alta", f"≥ {cortes['A']}",
              "Produção legislativa expressiva em vários temas da agenda, combinada com atuação em comissões "
              "temáticas (como titular ou na presidência) e em frentes parlamentares do setor.",
              "Relacionamento direto e contínuo, com reuniões institucionais da Presidência e da Diretoria da ABES. "
              "São candidatos naturais a relatorias e a presidências de comissões."),
        "B": ("Prioridade média", f"≥ {cortes['B']} e < {cortes['A']}",
              "Atuação consistente, em geral concentrada em poucos temas ou ancorada em comissões temáticas e, "
              "no Senado, em relatorias.",
              "Agenda periódica por tema e envio de notas técnicas. Acompanhar quem pode subir para a faixa A ao "
              "assumir postos nas comissões em 2027."),
        "C": ("Interesse pontual", f"≥ {cortes['C']} e < {cortes['B']}",
              "Uma ou poucas proposições na pauta, ou presença em comissão temática como suplente ou em "
              "frentes parlamentares.",
              "Comunicação segmentada por tema; acionar quando matéria do seu tema ou do seu estado estiver em "
              "pauta."),
        "D": ("Interesse baixo", f"> 0 e < {cortes['C']}",
              "Sinal fraco: normalmente apenas adesão a frentes parlamentares ou uma proposição isolada.",
              "Monitoramento e envio de materiais institucionais gerais; reavaliar ao longo da legislatura."),
    }
    linhas, cores = [], []
    for letra, (nome, faixa_iip, perfil, abordagem) in descricoes.items():
        grupo = h[h["prioridade"].str.startswith(letra)]
        cam, sen = int((grupo["casa"] == "Câmara").sum()), int((grupo["casa"] == "Senado").sum())
        linhas.append([f"{letra}\n{nome}", faixa_iip, f"{cam + sen}\n({cam} Câmara / {sen} Senado)",
                       perfil, abordagem])
        cores.append(COR_FAIXA[letra])
    _tabela(doc, ["Faixa", "IIP", "Parlamentares", "Perfil típico", "Abordagem recomendada"], linhas,
            [2.3, 1.7, 2.6, 4.7, 4.7], tamanho=8, cores_col0=cores, centro={1, 2})
    _para(doc, "")
    sem_sinal = int((hist["iip"].fillna(0) == 0).sum())
    _para(doc, (f"Fora das faixas: {sem_sinal} parlamentares com mandato atual não apresentaram nenhum sinal de "
                f"interesse na pauta (IIP = 0), e os {len(novos)} eleitos sem mandato anterior não são pontuados "
                f"pelo índice."), size=9, color=CINZA_RODAPE)

    # ── 4. Panorama ─────────────────────────────────────────────────────
    _heading1(doc, "4. Panorama das prioridades")
    img = _grafico_faixas(h)
    doc.add_picture(str(img), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    ordem = ["A — Prioridade alta", "B — Prioridade média", "C — Interesse pontual", "D — Interesse baixo"]
    linhas = []
    for faixa in ordem:
        cam = int(((h["casa"] == "Câmara") & (h["prioridade"] == faixa)).sum())
        sen = int(((h["casa"] == "Senado") & (h["prioridade"] == faixa)).sum())
        linhas.append([faixa, str(cam), str(sen), str(cam + sen)])
    linhas.append(["Total", str(int((h["casa"] == "Câmara").sum())), str(int((h["casa"] == "Senado").sum())),
                   str(len(h))])
    _tabela(doc, ["Faixa", "Câmara", "Senado", "Total"], linhas, [6.0, 2.5, 2.5, 2.5], tamanho=9,
            cores_col0=[COR_FAIXA[f[0]] for f in ordem] + [PRETO], centro={1, 2, 3})
    _para(doc, "")
    ab = pd.concat([fa, fb])
    partidos = Counter(ab["partido"]).most_common(8)
    _para(doc, (
        "Distribuição partidária das faixas A e B: "
        + ", ".join(f"{p} ({n})" for p, n in partidos)
        + ". O interesse pela pauta é transversal ao espectro político, o que favorece uma agenda "
          "institucional suprapartidária."
    ))

    # ── 5. Prioridade A ─────────────────────────────────────────────────
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _heading1(doc, f"5. Prioridade alta — faixa A ({len(fa)} parlamentares)")
    _para(doc, "Interlocutores de primeira ordem: combinam produção legislativa relevante na pauta com presença "
               "nas comissões temáticas.")
    linhas = []
    for pos, (_, r) in enumerate(fa.iterrows(), 1):
        linhas.append([f"{r['iip']:.0f}", f"{r['nome'].title()}\n{r['partido']}/{r['uf']} · {r['casa']}",
                       r["trajetoria"], _sinais_resumo(interesses[r["chave"]], nomes_temas)])
    _tabela(doc, ["IIP", "Parlamentar", "Trajetória", "Principais sinais de interesse"], linhas,
            [1.2, 4.3, 3.0, 7.5], tamanho=8, cores_col0=["2E7D32"] * len(linhas))

    # ── 6. Prioridade B ─────────────────────────────────────────────────
    _heading1(doc, f"6. Prioridade média — faixa B ({len(fb)} parlamentares)")
    linhas = [[f"{r['iip']:.0f}", r["nome"].title(), f"{r['partido']}/{r['uf']}", r["casa"], r["trajetoria"],
               str(int(r["n_proposicoes"] or 0)), sorted({o.split(' — ')[0] for o in interesses[r['chave']].orgaos})
               and ", ".join(sorted({o.split(' — ')[0] for o in interesses[r['chave']].orgaos})[:3]) or "—"]
              for _, r in fb.iterrows()]
    _tabela(doc, ["IIP", "Parlamentar", "Partido/UF", "Casa", "Trajetória", "Proposições", "Órgãos temáticos"],
            linhas, [1.1, 4.0, 2.4, 1.5, 3.0, 1.6, 2.4], tamanho=7.5,
            cores_col0=["7DBE5A"] * len(linhas), centro={3, 5})

    # ── 7. Por tema ─────────────────────────────────────────────────────
    _heading1(doc, "7. Referências por tema da agenda")
    _para(doc, (
        "O quadro a seguir cruza os temas da agenda ABES, descritos na seção 3, com a produção legislativa dos "
        "parlamentares que estarão no Congresso a partir de 2027. Para cada tema, são listados os cinco maiores "
        "autores de proposições entre 2023 e 2026 (entre parênteses, o número de proposições; no Senado, "
        "coautorias contam meio ponto). O quadro permite direcionar pautas específicas: por exemplo, levar a "
        "agenda de inteligência artificial aos parlamentares que mais legislaram sobre o tema, ainda que não "
        "estejam nas faixas de prioridade mais altas. Como no restante do relatório, a lista indica quem atua no "
        "tema, não a posição adotada."
    ))
    linhas = []
    for tema in agenda["temas"]:
        ranking = sorted(((i.autoria_por_tema.get(tema["id"], 0), k) for k, i in interesses.items()
                          if i and i.autoria_por_tema.get(tema["id"], 0) > 0), reverse=True)[:5]
        nomes = "; ".join(f"{por_chave[k].nome_urna.title()} ({por_chave[k].partido}/{por_chave[k].uf}, {n:g})"
                          for n, k in ranking) or "—"
        linhas.append([tema["nome"], nomes])
    _tabela(doc, ["Tema", "Principais autores"], linhas, [4.0, 12.0], tamanho=8)

    # ── 8. Recomendações ────────────────────────────────────────────────
    _heading1(doc, "8. Recomendações e próximos passos")
    recs = [
        ("Até a posse (nov/2026 – jan/2027): ", f"agendar apresentações institucionais da ABES com os {len(fa)} "
         "parlamentares da faixa A, priorizando quem tende a disputar presidências e relatorias nas comissões "
         "temáticas."),
        ("Instalação das comissões (fev – mar/2027): ", "acompanhar a escolha de presidentes e membros da CCTI e "
         "CCOM (Câmara) e da CCT e CCDD (Senado), cruzando com as faixas A e B para antecipar relatorias de "
         "matérias prioritárias (PL 2338/2023, PL 4752/2025, regulamentação da reforma tributária)."),
        ("Qualificar posição: ", f"aplicar às faixas A e B ({len(fa) + len(fb)} nomes) a análise de afinidade "
         "(nota de −2 a +2 por tema, com fontes), para separar aliados de interlocutores críticos antes das "
         "reuniões."),
        ("Mapear os novos parlamentares: ", f"os {len(novos)} eleitos sem mandato anterior devem ser mapeados por "
         "propostas de campanha, trajetória profissional e presença digital."),
        ("Validar a régua: ", "submeter à Diretoria de RIG os temas, pesos e palavras-chave da agenda ABES "
         "utilizados no índice (eleicoes/agenda_abes.json) e recalcular após eventuais ajustes."),
    ]
    for neg, txt in recs:
        _bullet(doc, txt, negrito=neg)

    # ── 9. Limitações ───────────────────────────────────────────────────
    _heading1(doc, "9. Limitações")
    for txt in [
        "O IIP mede interesse/engajamento, não alinhamento às posições da ABES.",
        "A identificação de temas usa palavras-chave nas ementas; pode haver falsos positivos e omissões. "
        "As evidências de cada parlamentar estão na planilha que acompanha este relatório.",
        "Votos nominais, discursos e relatorias na Câmara não foram considerados nesta versão.",
        "Senadores que assumirem governos estaduais em 2027 serão substituídos por suplentes, "
        "que não estão avaliados aqui.",
    ]:
        _bullet(doc, txt)
    _para(doc, "Anexo: planilha prioridade_agenda_ABES.xlsx — lista completa hierarquizada, evidências por "
               "parlamentar e metodologia.", size=9, color=CINZA_RODAPE, space_before=6)

    # ── 10. Nota de versão ──────────────────────────────────────────────
    _heading1(doc, "10. Nota de versão")
    _para(doc, f"Esta edição ({hoje.strftime('%d/%m/%Y')}) atualiza a versão de 05/10/2026 com as seguintes "
               f"alterações:")
    for neg, txt in [
        ("Dados atualizados: ", f"composição e índice recalculados com dados coletados em {comp.data_coleta}. "
         "No Senado, Ana Paula Lobato (PSB/MA) reassumiu em 06/10/2026 o mandato até 2031, no lugar da "
         "suplente Lourdinha Pereira, e passa a integrar a faixa D. Com isso, o total de parlamentares com "
         "interesse prévio na pauta passou de 351 para 352; a única outra variação foi o IIP de Flávia Morais "
         "(MDB/GO), de 37,5 para 39,5, que permanece na faixa B."),
        ("Reprocessamento em Pernambuco: ", "o TSE sinalizou reprocessamento do resultado de deputados federais "
         "de PE. Foi mantida a lista validada em 05/10/2026, que poderá mudar após a conclusão do processo."),
        ("Panorama das prioridades: ", "o gráfico passou a rotular todos os segmentos — o parlamentar do Senado "
         "na faixa A (Astronauta Marcos Pontes) não aparecia — e foi incluída tabela de conferência faixa × casa."),
        ("Faixas de prioridade: ", "descrição de cada faixa, com intervalo do IIP, perfil típico e abordagem "
         "recomendada."),
        ("Temas da agenda ABES: ", "explicação de como os temas, pesos e palavras-chave são usados na "
         "categorização, com quadro de escopo de cada tema."),
    ]:
        _bullet(doc, txt, negrito=neg)

    caminho = caminho or DATA_DIR / f"relatorio_prioridade_agenda_ABES_{hoje.isoformat()}.docx"
    # Títulos nunca ficam sozinhos no pé da página
    for par in doc.paragraphs:
        if par.runs and par.runs[0].bold and par.runs[0].font.size and par.runs[0].font.size >= Pt(12):
            par.paragraph_format.keep_with_next = True
    _ordenar_xml(doc)
    doc.save(caminho)
    img.unlink(missing_ok=True)
    return caminho


if __name__ == "__main__":
    print(gerar())
